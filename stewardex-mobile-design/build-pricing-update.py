# -*- coding: utf-8 -*-
"""Generate the client-facing Stewardex Plans & Pricing document as a styled .docx.
Four plans: Starter, Professional, Organisational, Enterprise. Australian English."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL   = RGBColor(0x1F, 0x8A, 0x7E)
NAVY   = RGBColor(0x0B, 0x25, 0x45)
INK    = RGBColor(0x22, 0x2A, 0x33)
MUTED  = RGBColor(0x5A, 0x66, 0x75)
TEAL_HEX, LIGHT_HEX, ZEBRA_HEX, NEW_HEX = "1F8A7E","EAF3F1","F4F7F6","FBF3E6"

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.13
for s in doc.sections:
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.5); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto"); shd.set(qn("w:fill"),hexcolor); tcPr.append(shd)

def cell_text(cell, text, bold=False, color=INK, size=9, align=None, white=False):
    cell.text = ""; p = cell.paragraphs[0]
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    for i, line in enumerate(str(text).split("\n")):
        r = p.add_run(("\n" if i else "") + line); r.bold = bold; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF) if white else color
    return p

def no_borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top","left","bottom","right","insideH","insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"),"none"); b.append(x)
    tblPr.append(b)

def heading(num, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{num}.  {text}"); r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = NAVY
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"),"single"); bot.set(qn("w:sz"),"12"); bot.set(qn("w:space"),"3"); bot.set(qn("w:color"),TEAL_HEX)
    pBdr.append(bot); pPr.append(pBdr); return p

def subhead(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = TEAL; return p

def body(text, italic=False, color=INK, size=10.5, space=5, bold=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text); r.italic = italic; r.bold = bold; r.font.color.rgb = color; r.font.size = Pt(size); return p

def bullet(text, lead=None):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.1
    if lead:
        r = p.add_run(lead + "  "); r.bold = True
    r2 = p.add_run(text); r2.font.size = Pt(10.5); return p

def make_table(headers, rows, widths=None, fontsize=9, zebra=True, highlight_rows=None, center_cols=None):
    highlight_rows = highlight_rows or set(); center_cols = center_cols or set()
    t = doc.add_table(rows=1+len(rows), cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        shade(t.cell(0,i), TEAL_HEX)
        cell_text(t.cell(0,i), h, bold=True, white=True, size=fontsize,
                  align=(WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else None))
    for ri,row in enumerate(rows, start=1):
        for ci,val in enumerate(row):
            c = t.cell(ri,ci)
            if ri in highlight_rows: shade(c, NEW_HEX)
            elif zebra and ri % 2 == 0: shade(c, ZEBRA_HEX)
            cell_text(c, val, size=fontsize, bold=(ci==0),
                      align=(WD_ALIGN_PARAGRAPH.CENTER if ci in center_cols else None))
    if widths:
        for ri in range(len(rows)+1):
            for ci,w in enumerate(widths):
                t.cell(ri,ci).width = w
    return t

# ============================================================ TITLE
t = doc.add_paragraph(); t.paragraph_format.space_after = Pt(0)
r = t.add_run("STEWARDEX — PLANS & PRICING"); r.bold = True; r.font.size = Pt(23); r.font.color.rgb = NAVY
st = doc.add_paragraph(); st.paragraph_format.space_after = Pt(2)
r = st.add_run("Updated plans, pricing and new capabilities"); r.font.size = Pt(12); r.font.color.rgb = TEAL; r.bold = True
meta = doc.add_table(rows=2, cols=4)
for i,(k,v) in enumerate([("Prepared for","[Client name]"),("Prepared by","[Your company]"),
                          ("Currency","AUD"),("Date","19 June 2026")]):
    cell_text(meta.cell(0,i), k.upper(), bold=True, color=MUTED, size=8)
    cell_text(meta.cell(1,i), v, bold=True, color=NAVY, size=9.5)
no_borders(meta)
rule = doc.add_paragraph(); rule.paragraph_format.space_before = Pt(4); rule.paragraph_format.space_after = Pt(2)
pPr = rule._p.get_or_add_pPr(); pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
bot.set(qn("w:val"),"single"); bot.set(qn("w:sz"),"18"); bot.set(qn("w:space"),"1"); bot.set(qn("w:color"),TEAL_HEX)
pBdr.append(bot); pPr.append(pBdr)

# ============================================================ 1. OVERVIEW
heading(1, "Overview")
body("Stewardex has grown considerably since the plans were first set. The platform now spans governance and "
     "compliance, finance, fundraising and donor management, project and grant delivery, people and volunteers, "
     "marketing, team communication, and a companion mobile app. This document sets out the updated four-plan "
     "line-up and current pricing, explains the new capabilities and where they sit, and shows what each plan "
     "includes.")
body("All prices are in Australian dollars (AUD) and exclude GST. Annual billing is offered at a 10% saving versus "
     "paying monthly.", space=4)

# ============================================================ 2. PLAN LINE-UP & PRICING
heading(2, "Plan line-up & pricing")
make_table(
    ["Plan","Monthly","Annual (save 10%)","Best for","Staff seats","Approvals / mo","Storage"],
    [["Starter","$399","$4,309","Small charities ($50K–$1M)","5","50","20 GB"],
     ["Professional","$1,499","$16,189","Medium charities ($1M–$5M)","25","750","100 GB"],
     ["Organisational","$5,999","$64,789","Larger charities ($5M–$10M)","Unlimited","1,500","Unlimited"],
     ["Enterprise","Contact us","Contact us","Large organisations & groups ($10M+)","Unlimited","Unlimited","Unlimited"]],
    widths=[Cm(2.7),Cm(1.9),Cm(2.3),Cm(4.4),Cm(1.9),Cm(1.9),Cm(1.9)], fontsize=9,
    center_cols={1,2,4,5,6})
body("“Approvals / month” is the included allowance of completed approval workflows. Board and committee seats are "
     "included on every plan (9 on Starter; unlimited from Professional upward). Annual plans waive the one-off "
     "setup fee. Enterprise is tailored to the size and structure of the largest organisations and federations — "
     "pricing is provided on request.", italic=True, color=MUTED, size=9.5, space=2)

# ============================================================ 3. NEW CAPABILITIES
heading(3, "New capabilities on the platform")
body("The following capabilities have been added since the original plans and are reflected in the comparison in "
     "section 4.", space=4)

subhead("Fundraising & engagement")
bullet("manage donors and donations, donation milestones, donation boxes, and refunds in one place.", "Fundraising & Donor Management —")
bullet("register funded projects, track delivery against funding agreements, and manage over-budget change requests with partners.", "Project & Grant Delivery —")
bullet("plan campaigns and route marketing spend through approval before it goes out.", "Marketing & Social Campaigns —")
bullet("recruit, induct and manage volunteers alongside your people records.", "Volunteer Management —")

subhead("Access anywhere & collaboration")
bullet("a companion iOS & Android app for approvals, registers, meetings, training and expenses on the go, with push notifications.", "Mobile App —")
bullet("secure in-app messaging across channels, with mentions and threads, so conversations stay next to the work.", "Team Chat —")

subhead("Everyday tools (included on every plan)")
bullet("a shared organisation calendar with reminders for due and expiring items.", "Calendar & Reminders —")
bullet("configurable inquiry records that move through their own approval workflow.", "Inquiry Register —")
bullet("raise and track support requests with our team from inside the platform.", "In-app Support —")

# ============================================================ 4. WHAT EACH PLAN INCLUDES
heading(4, "What each plan includes")
body("A ✓ means the capability is included. “Add-on” means it can be added to that plan for an extra monthly fee. "
     "Enterprise includes everything in Organisational, tailored to your structure with the highest level of "
     "onboarding and support.", space=4)
make_table(
    ["Capability","Starter","Professional","Organisational","Enterprise"],
    [["Charity administration, board pack & meetings","✓","✓","✓","✓"],
     ["Policies, risk, complaints & incidents registers","✓","✓","✓","✓"],
     ["Compliance checklists & AIS lodgement","✓","✓","✓","✓"],
     ["Audit trail, two-step login & role permissions","✓","✓","✓","✓"],
     ["External auditor access","✓","✓","✓","✓"],
     ["Inquiry register, calendar & in-app support","✓","✓","✓","✓"],
     ["Fundraising & Donor Management","—","✓","✓","✓"],
     ["Project & Grant Delivery","—","✓","✓","✓"],
     ["Marketing & Social Campaigns","—","✓","✓","✓"],
     ["Volunteer Management","—","✓","✓","✓"],
     ["Mobile App (iOS & Android)","Add-on","✓","✓","✓"],
     ["Team Chat & channels","Add-on","✓","✓","✓"],
     ["Finance suite — expenses, invoices, budgets, BAS, statements","—","✓","✓","✓"],
     ["Cash handling & sweep funds","—","✓","✓","✓"],
     ["Partner & donor vetting","—","✓","✓","✓"],
     ["People & HR · electronic signing for minutes","—","✓","✓","✓"],
     ["AI Compliance Assistant","—","✓","✓","✓"],
     ["Supplier, systems (IT) & insurance registers","—","✓","✓","✓"],
     ["Multi-entity management & custom workflow builder","—","—","✓","✓"],
     ["Business continuity plan & shared credential vault","—","—","✓","✓"],
     ["Regulatory change monitoring","—","—","✓","✓"],
     ["Single sign-on (SSO/SCIM), developer API & webhooks","—","—","✓","✓"],
     ["White-label branding & choice of data region","—","—","✓","✓"],
     ["ISO 27001 / SOC 2 evidence pack","—","—","✓","✓"],
     ["Dedicated success manager & 99.9% uptime guarantee","—","—","✓","✓"],
     ["Tailored onboarding & configuration to your structure","—","—","—","✓"]],
    widths=[Cm(8.0),Cm(2.25),Cm(2.25),Cm(2.25),Cm(2.25)], fontsize=8.5, center_cols={1,2,3,4},
    highlight_rows={7,8,9,10,11,12})
body("Newer fundraising, delivery, marketing and engagement modules are included from Professional upward, so the "
     "richer toolset sits with the charities that use it most.", italic=True, color=MUTED, size=9.5, space=2)

# ============================================================ 5. ADD-ONS & EXTRAS
heading(5, "Add-ons & extras")
bullet("available to Starter customers as an add-on (suggested $99 / month) so the smallest teams can work on the "
       "go without moving up a plan. Included at no extra cost from Professional upward.", "Mobile App —")
bullet("purchase individual, ready-to-use policy templates as a one-off, on any plan — a fast way to fill a gap in "
       "your policy library.", "Policy Template Marketplace —")

# ============================================================ 6. SUMMARY
heading(6, "Summary")
body("The updated line-up keeps Starter as an accessible entry point for small charities, gives Professional the "
     "full fundraising, delivery, finance and people toolset for growing organisations, and reserves Organisational "
     "and Enterprise for larger charities and groups that need multi-entity management, advanced security and "
     "tailored support. The new mobile app, team chat, fundraising, delivery and marketing capabilities add real "
     "value across the range and give customers a clear path to grow with Stewardex.")

foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.CENTER; foot.paragraph_format.space_before = Pt(10)
r = foot.add_run("Stewardex · Plans & Pricing · Prices in AUD, excl. GST"); r.font.size = Pt(8); r.font.color.rgb = MUTED

out = "Stewardex-Plans-and-Pricing-v2.docx"
doc.save(out)
print("Saved", out)
