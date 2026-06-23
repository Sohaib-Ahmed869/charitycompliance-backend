# -*- coding: utf-8 -*-
"""Generate the Stewardex Mobile project proposal as a styled .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL   = RGBColor(0x1F, 0x8A, 0x7E)   # brand accent
NAVY   = RGBColor(0x0B, 0x25, 0x45)   # headings
INK    = RGBColor(0x22, 0x2A, 0x33)   # body
MUTED  = RGBColor(0x5A, 0x66, 0x75)
TEAL_HEX = "1F8A7E"
NAVY_HEX = "0B2545"
LIGHT_HEX = "EAF3F1"
ZEBRA_HEX = "F4F7F6"

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12

# tighter margins for a clean 3-4 pager
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=INK, size=10, align=None, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF) if white else color
    return p

def heading(num, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{num}.  {text}")
    r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = NAVY
    # thin accent underline
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"12")
    bottom.set(qn("w:space"),"3"); bottom.set(qn("w:color"),TEAL_HEX)
    pBdr.append(bottom); pPr.append(pBdr)
    return p

def subhead(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = TEAL
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.1
    if bold_lead:
        r = p.add_run(bold_lead + "  "); r.bold = True; r.font.color.rgb = INK
    r2 = p.add_run(text); r2.font.size = Pt(10.5)
    return p

def body(text, italic=False, color=INK, size=10.5, space=5):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text); r.italic = italic; r.font.color.rgb = color; r.font.size = Pt(size)
    return p

# =====================================================================
# TITLE BLOCK
# =====================================================================
t = doc.add_paragraph(); t.paragraph_format.space_after = Pt(0)
r = t.add_run("STEWARDEX MOBILE"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
st = doc.add_paragraph(); st.paragraph_format.space_after = Pt(2)
r = st.add_run("Project Proposal — Phase 1 Mobile Application (iOS & Android)")
r.font.size = Pt(12); r.font.color.rgb = TEAL; r.bold = True

# meta table (no borders)
meta = doc.add_table(rows=2, cols=4); meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = True
labels = [("Prepared for","[Client name]"),("Prepared by","[Your company]"),
          ("Date","19 June 2026"),("Valid until","19 July 2026")]
for i,(k,v) in enumerate(labels):
    set_cell_text(meta.cell(0,i), k.upper(), bold=True, color=MUTED, size=8)
    set_cell_text(meta.cell(1,i), v, bold=True, color=NAVY, size=10)
for row in meta.rows:
    for c in row.cells:
        c._tc.get_or_add_tcPr().append(OxmlElement("w:tcMar"))
# remove table borders
tblPr = meta._tbl.tblPr
borders = OxmlElement("w:tblBorders")
for edge in ("top","left","bottom","right","insideH","insideV"):
    e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"none"); borders.append(e)
tblPr.append(borders)

# accent rule
rule = doc.add_paragraph(); rule.paragraph_format.space_before = Pt(4); rule.paragraph_format.space_after = Pt(2)
pPr = rule._p.get_or_add_pPr(); pBdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"18"); bottom.set(qn("w:space"),"1"); bottom.set(qn("w:color"),TEAL_HEX)
pBdr.append(bottom); pPr.append(pBdr)

# =====================================================================
# 1. OVERVIEW
# =====================================================================
heading(1, "Overview")
body("Stewardex Mobile is a native-feel, cross-platform companion application (iOS & Android) for the "
     "Stewardex governance and compliance suite. It puts the day-to-day, time-sensitive work — approvals, "
     "workflows, registers, meetings, chat and training — in the pocket of every team member, so that nothing "
     "that requires a person stalls while they are away from their desk.")
body("The application runs on the same server and infrastructure as the existing Stewardex web platform. "
     "Every action performed on mobile is recorded with the identical workflow processing, audit trail and "
     "approval logic as the web app — the mobile experience aids and fast-tracks the existing software; it does "
     "not replace or fork it.")
body("This proposal covers Phase 1: the scope, a four-week delivery timeline (three weeks of development and a "
     "dedicated week of UAT & QA), and the commercials.")

# =====================================================================
# 2. OBJECTIVES
# =====================================================================
heading(2, "Objectives")
bullet("so that no approval, decision or resolution becomes a hidden bottleneck.", "Surface every blocker —")
bullet("create, triage, route/hand-off and resolve register items directly from mobile.", "Action the registers —")
bullet("keep workflow recording, audit trail and notifications identical to the web app.", "Preserve parity —")
bullet("administrative configuration stays on web; mobile gives the right people the right action.", "Respect permissions —")

# =====================================================================
# 3. SCOPE
# =====================================================================
heading(3, "Scope")
subhead("In scope — Phase 1")
bullet("secure sign-in, two-factor / OTP, Face ID unlock and role-based access.", "Authentication & security:")
bullet("today's tasks, pending approvals, quick actions and an iOS-style notification centre.", "Home / My Tasks:")
bullet("real-time alerts for approvals, mentions, meetings, policies, risks and complaints, each deep-linking "
       "to the relevant record.", "Push notifications:")
bullet("month, week, list and 'expiring soon' views; add meetings, events and personal tasks; reminders for "
       "due and expiring items.", "Calendar:")
bullet("the full approval inbox with approval-chain visibility — approve, reject, escalate, flag a COI, "
       "raise/route a risk, or delegate. All workflow types surface on mobile.", "Approvals & workflows:")
bullet("channel and direct messaging, mentions and threads (joining existing channels; creating new channels "
       "remains on web).", "Communication (chat):")
bullet("read records and create new items, plus the workflow handling for each —", "Registers — view, create & resolve:")
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.space_after = Pt(2)
for lead, txt in [
    ("Risk:", "heat-map & records (read); raise a risk; admin triage — set priority, assign owner and route / hand off to a head of department; escalate; add treatment; resolve."),
    ("Conflicts of Interest:", "declare a conflict; review chain; route to board; add mitigation; approve or reject."),
    ("Complaints:", "log a complaint; assign a handler; escalate; respond; resolve and close."),
    ("Policies:", "read and acknowledge policies (uploading new policies remains on web)."),
]:
    sp = doc.add_paragraph(style="List Bullet 2"); sp.paragraph_format.space_after = Pt(1)
    r = sp.add_run(lead + " "); r.bold = True; r.font.size = Pt(10)
    r2 = sp.add_run(txt); r2.font.size = Pt(10)
bullet("upcoming & past meetings, RSVP, agenda, minutes & notes, comments, and add a meeting.", "Meetings:")
bullet("take and watch assigned training and videos, resume in place, with completion and watch-time recorded "
       "to the user's training record.", "Training:")
bullet("submit a claim against vetted suppliers, allocate to a project budget and attach an invoice.", "Expenses:")
bullet("view plan, usage and invoices, upgrade, and manage payment via Stripe.", "Billing:")
bullet("administrative actions (allocation, decision, resolution) are gated by role; non-privileged users "
       "receive read access plus their own permitted action. Light & dark mode throughout.", "Role-aware permissions:")

subhead("Out of scope — Phase 1")
bullet("Configuring approval workflows and permission matrices.")
bullet("Full register administration / set-up and supplier–partner vetting (vetting status is read-only on mobile).")
bullet("Organisation onboarding and set-up wizards.")
bullet("Uploading new policies or legal documents.")
bullet("Advanced reporting — fiscal reports and grant-owner management.")
bullet("The Project Delivery register as a dedicated screen (its workflows still surface on mobile; a full "
       "register view is available as an optional add-on).")

# =====================================================================
# 4. TIMELINE
# =====================================================================
heading(4, "Timeline — 4 weeks (3 development + 1 UAT & QA)")
tl = doc.add_table(rows=5, cols=3); tl.style = "Table Grid"; tl.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Cm(2.2), Cm(5.2), Cm(9.0)]
hdr = ["Week", "Stage", "Key activities"]
for i,h in enumerate(hdr):
    shade(tl.cell(0,i), TEAL_HEX); set_cell_text(tl.cell(0,i), h, bold=True, white=True, size=10)
rows = [
    ("Week 1", "UI build & integration setup",
     "Finalise screens from the approved prototype; design system & navigation; authentication; "
     "notification scaffolding; backend API integration kickoff."),
    ("Week 2", "Core development I",
     "Home / My Tasks, calendar, approvals & workflows, push notifications & deep-linking, chat."),
    ("Week 3", "Core development II",
     "Registers (risk, COI, complaints — create + workflow handling), policies, meetings, training, "
     "expenses, billing and role-based gating. Feature-complete build."),
    ("Week 4", "UAT, QA & release",
     "Full QA across iOS & Android, client UAT, bug-fix cycles, app-store submission preparation and sign-off."),
]
for ri,(w,s,a) in enumerate(rows, start=1):
    if ri % 2 == 0:
        for ci in range(3): shade(tl.cell(ri,ci), ZEBRA_HEX)
    set_cell_text(tl.cell(ri,0), w, bold=True, color=NAVY, size=10)
    set_cell_text(tl.cell(ri,1), s, bold=True, color=INK, size=10)
    set_cell_text(tl.cell(ri,2), a, size=9.5)
body("Milestones:  Kickoff (start of Week 1)  •  Feature-complete (end of Week 3)  •  UAT commences (start of "
     "Week 4)  •  Go-live / acceptance (end of Week 4).", italic=True, color=MUTED, size=9.5, space=2)

# =====================================================================
# 5. APP STORE APPROVAL
# =====================================================================
heading(5, "App store submission & approval (Apple & Google)")
body("Publishing to the Apple App Store and Google Play requires a review and approval by Apple and Google. "
     "This review is performed by the stores themselves and is outside our control — and therefore it sits "
     "separate to, and runs after, the four-week delivery timeline above. We prepare and lodge the store "
     "submissions during Week 4 (UAT & QA); the review clock then starts once each store accepts the build.")
sub = doc.add_table(rows=3, cols=3); sub.style = "Table Grid"; sub.alignment = WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(["Store", "Typical review window", "Notes"]):
    shade(sub.cell(0,i), TEAL_HEX); set_cell_text(sub.cell(0,i), h, bold=True, white=True, size=10)
sub_rows = [
    ("Apple App Store", "1–3 business days",
     "First-ever submission can take longer; rejections add a further review cycle."),
    ("Google Play", "A few hours – 3 days",
     "New developer accounts may be held for extended review (up to ~7 days)."),
]
for ri,(s,w,n) in enumerate(sub_rows, start=1):
    if ri % 2 == 0:
        for ci in range(3): shade(sub.cell(ri,ci), ZEBRA_HEX)
    set_cell_text(sub.cell(ri,0), s, bold=True, color=NAVY, size=10)
    set_cell_text(sub.cell(ri,1), w, size=10)
    set_cell_text(sub.cell(ri,2), n, size=9.5)
body("Indicative allowance for store approval: approximately 1 week after go-live, separate to the development "
     "timeline. Prerequisites — supplied by the client before submission — are an active Apple Developer Program "
     "account and a Google Play Console account, together with final store-listing assets (icon, screenshots, "
     "description and privacy policy). We manage the submission, respond to reviewer queries and handle one "
     "resubmission per store if required.", italic=True, color=MUTED, size=9.5, space=2)

# =====================================================================
# 6. COMMERCIALS
# =====================================================================
heading(6, "Commercials")
body("Engagement model: fixed-price, Phase 1. All amounts are in Australian dollars (AUD) and exclude GST.",
     color=INK, size=10.5, space=4)
body("Indicative pricing — to be confirmed before issuing to the client.", italic=True, color=RGBColor(0xB0,0x4A,0x1A), size=9.5, space=4)

fee = doc.add_table(rows=5, cols=2); fee.style = "Table Grid"; fee.alignment = WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(["Deliverable", "Fee (AUD, ex GST)"]):
    shade(fee.cell(0,i), TEAL_HEX); set_cell_text(fee.cell(0,i), h, bold=True, white=True, size=10,
        align=(WD_ALIGN_PARAGRAPH.RIGHT if i==1 else None))
fee_rows = [
    ("UI build & integration setup (Week 1)", "$6,000"),
    ("Core development (Weeks 2–3)", "$12,000"),
    ("UAT, QA & release (Week 4)", "$6,000"),
]
for ri,(d,f) in enumerate(fee_rows, start=1):
    if ri % 2 == 0:
        for ci in range(2): shade(fee.cell(ri,ci), ZEBRA_HEX)
    set_cell_text(fee.cell(ri,0), d, size=10)
    set_cell_text(fee.cell(ri,1), f, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
shade(fee.cell(4,0), LIGHT_HEX); shade(fee.cell(4,1), LIGHT_HEX)
set_cell_text(fee.cell(4,0), "Total fixed fee (ex GST)", bold=True, color=NAVY, size=10.5)
set_cell_text(fee.cell(4,1), "$24,000 AUD", bold=True, color=NAVY, size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

subhead("Payment schedule")
bullet("50% on commencement — $12,000 AUD.")
bullet("30% at UAT commencement (start of Week 4) — $7,200 AUD.")
bullet("20% on acceptance / go-live — $4,800 AUD.")

subhead("Included")
bullet("Project management, design polish from the approved prototype, and a single cross-platform build "
       "(iOS + Android).")
bullet("Integration with the existing Stewardex backend, QA across both platforms, and UAT support.")
bullet("One submission per app store and a 2-week post-launch warranty for defect fixes.")

subhead("Optional add-ons (quoted separately)")
bullet("Project Delivery register (full mobile view).")
bullet("Advanced reporting / fiscal dashboards on mobile.")
bullet("Additional feature iterations or UAT cycles, charged at an agreed day rate.")

# =====================================================================
# 7. ASSUMPTIONS & DEPENDENCIES
# =====================================================================
heading(7, "Assumptions & dependencies")
bullet("The interactive prototype is the approved design baseline for Phase 1.")
bullet("Existing backend services are available and extended as needed; mobile reuses the current infrastructure.")
bullet("The client provides timely UAT feedback within the Week 4 window, and the Apple Developer and Google "
       "Play accounts required for store submission.")
bullet("Content such as training videos and policy documents is supplied by the client.")
bullet("Changes beyond the stated scope are handled via a written change request and may affect timeline and fee.")

# =====================================================================
# 8. NEXT STEPS
# =====================================================================
heading(8, "Next steps")
bullet("Confirm scope and commercials, and countersign this proposal.")
bullet("Settle the commencement invoice to schedule the kickoff.")
bullet("Begin Week 1 — UI build & integration setup, targeting go-live at the end of Week 4.")

# acceptance line
doc.add_paragraph().paragraph_format.space_after = Pt(6)
acc = doc.add_table(rows=1, cols=2); acc.alignment = WD_TABLE_ALIGNMENT.LEFT
for i,lab in enumerate(["Accepted for [Client name]", "Accepted for [Your company]"]):
    set_cell_text(acc.cell(0,i), "\n\n_______________________________\n" + lab + "\nName / Signature / Date",
                  color=MUTED, size=9)
tblPr = acc._tbl.tblPr; borders = OxmlElement("w:tblBorders")
for edge in ("top","left","bottom","right","insideH","insideV"):
    e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"none"); borders.append(e)
tblPr.append(borders)

foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.space_before = Pt(10)
r = foot.add_run("Stewardex Mobile · Phase 1 Proposal · Confidential")
r.font.size = Pt(8); r.font.color.rgb = MUTED

out = "Stewardex-Mobile-Proposal.docx"
doc.save(out)
print("Saved", out)
